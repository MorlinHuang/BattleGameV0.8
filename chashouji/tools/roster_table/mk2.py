import json,os,re,openpyxl
from openpyxl.styles import Font,Alignment,PatternFill,Border,Side
from openpyxl.drawing.image import Image as XImg
EK={'ride':'骑行/滑行进场','roll':'翻滚进场','walk':'走进来','appear':'原地显现','swing':'荡秋千进场','rope':'躺绳上荡入','fly':'飞入/降下','drop':'倒挂落下','leap':'跃起落下','slide':'滑入','creep':'匍匐爬入','crawl':'爬进来','dive':'飞扑滑入','dash':'冲刺进场','spring':'从屏幕边弹进来'}
AK={'jet':'连续喷射','rush':'突进连击（残影）','throw':'投掷飞行道具','whip':'甩鞭/绸抽打','slash':'剑光斩','beam':'光束/气浪','camera':'拍照闪光','punch':'伸缩捅击','spray':'喷雾/喷酒'}
RC={'pepper':'喷雾雾团','star':'金星','debris':'碎屑','thud':'闷击冲击圈','petal':'花瓣','water':'水花','rouge':'红粉/口红印','foxfire':'青蓝狐火','crescent':'金月牙+音符','hollow':'紫光球','splash':'溅水/酒花','feather':'羽毛纸屑（软击）','melon':'西瓜爆开'}
SL={'地面':'后排地面','上方':'上方','地板':'前景地板'}
CFG={'bestie':dict(P='G',title='闺蜜30人',dir='v14/trio/bestie',btn='闺蜜③',tgt='他',hf='F4C2D7',var='bestie'),
     'buddy':dict(P='B',title='哥们30人',dir='v14/trio/buddy',btn='哥们③',tgt='她',hf='BDD7EE',var='buddy')}
URL='http://1.14.252.30:40235/index.html?'
wb=openpyxl.Workbook();wb.remove(wb.active)
thin=Side(style='thin',color='CCCCCC');tf={'后排地面':'FFF2CC','上方':'DDEBF7','前景地板':'E2EFDA'}
for side in ['bestie','buddy']:
  c=CFG[side];d=json.load(open(f'/tmp/{side}.json'))
  src=open(f'web/trio_{side}.js').read();grp_ids={}
  for m in re.finditer(r"\{ name: '(.+?)', ground: '([GB]\d+)', top: '([GB]\d+)', floor: '([GB]\d+)' \}",src):
    for g in m.groups()[1:]: grp_ids[g]=f"{m.group(2)}.{m.group(3)}.{m.group(4)}"
  dirs={x.split('_')[0]:x for x in os.listdir(c['dir']) if re.match(c['P']+r'\d+_',x)}
  rows=[]
  for x in d:
    i=x['id'];s=x.get('sign') or {};ek=x.get('enterKind') or 'swing'
    rows.append({'编号':i,'形象（待机帧）':'','槽位':SL.get(x['slot'],x['slot']),'分组':x['group'],'角色（名单名）':x['name'],'年代':x['era'],
     '实际效果概述':x['head'],
     '进场':f"{EK.get(ek,ek)}（帧：{x.get('enterSeq') or 'idle'}）",
     '待机':f"帧 {x['idle'] or 'idle'}，呼吸起伏",
     '出手':f"{AK.get(x['atkKind'],x['atkKind'])}（帧：{x['atkSeq']}）",
     '出手道具/飞行物':(x['item']+'：' if x['item'] else '')+(x['props'] or '程序绘制（引擎画）'),
     '挂件/场景层':x['parts'] or '—',
     '命中特效':f"{x['recipe']}（{RC.get(x['recipe'],'')}）"+(f"，命中后 {x['onHit']}" if x['onHit'] else ''),
     '站位 at [x,y,缩放]':x['at'],'角色图集（web/assets/trio/）':x['sheet'],'图集帧名':x['frames'],
     '素材源目录':(c['dir']+'/'+dirs[i]) if i in dirs else '—（引擎维护）',
     '签字':f"{s.get('st','')}·{s.get('batch','')}",'已知小问题（非阻断）':s.get('note') or '—',
     '名单原设计·进场':x['rEnter'],'名单原设计·出手':x['rAtk'],
     f"查看链接（打开后点「{c['btn']}」）":URL+side+'='+grp_ids.get(i,i),'你的调优建议':''})
  ws=wb.create_sheet(c['title']);cols=list(rows[0].keys());ws.append(cols)
  for r in rows: ws.append([r[k] for k in cols])
  W={'编号':6,'形象（待机帧）':23,'槽位':9,'分组':8,'角色（名单名）':22,'年代':5,'实际效果概述':48,'进场':30,'待机':16,'出手':30,'出手道具/飞行物':26,'挂件/场景层':24,'命中特效':20,'站位 at [x,y,缩放]':16,'角色图集（web/assets/trio/）':22,'图集帧名':30,'素材源目录':26,'签字':14,'已知小问题（非阻断）':30,'名单原设计·进场':26,'名单原设计·出手':26,'你的调优建议':40}
  for j,k in enumerate(cols,1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width=W.get(k,40)
    h=ws.cell(1,j);h.font=Font(bold=True);h.fill=PatternFill('solid',fgColor=c['hf'])
  for row in ws.iter_rows():
    for cell in row: cell.alignment=Alignment(wrap_text=True,vertical='top');cell.border=Border(top=thin,bottom=thin,left=thin,right=thin)
  for k,r in enumerate(rows,2):
    im=XImg(f"/tmp/thumbs/{r['编号']}.png");im.width=im.height=160;ws.add_image(im,f'B{k}');ws.row_dimensions[k].height=125
    ws.cell(k,3).fill=PatternFill('solid',fgColor=tf.get(r['槽位'],'FFFFFF'));ws.cell(k,len(cols)).fill=PatternFill('solid',fgColor='FFFDE7')
  ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
ws2=wb.create_sheet('说明')
for t in ['两张表：闺蜜30人（G1~G30，左边，打男生）、哥们30人（B1~B30，右边，打女生）。',
 '槽位：后排地面 1~10，上方 11~20，前景地板 21~30；一次送礼三个槽位各随机抽 1 人。',
 '查看某人：打开「查看链接」（固定召出她/他所在那一组三人），再点页面下方「闺蜜③」或「哥们③」送礼；链接里的三个编号可以换成任意三人（每个槽位各一个）。',
 '帧名含义：idle 待机；wind/raise 蓄力；throw/fire/hit 出手；follow 收势；其余是进场帧（walk/ride/leap/land 等）。',
 '形象列是线上图集里的待机帧；挂件和场景层（无人机、月亮、墙沿、飞剑等）不在缩略图里，要在游戏里看。',
 '命中特效是引擎配方名，括号里是画面效果。图集在 /home/op/chashouji/web/assets/trio/，源素材在「素材源目录」。',
 '「你的调优建议」列留给你填，填完发回给我，我按编号分派给美术/引擎。']:
  ws2.append([t])
ws2.column_dimensions['A'].width=120
out='docs/三人组60人角色对照表.xlsx';wb.save(out);print(out)
